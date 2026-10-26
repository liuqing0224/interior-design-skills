#!/usr/bin/env python3
"""Original 96 sqm presentation regression; use verify_2d.py for generic projects."""
import argparse
from collections import Counter
import hashlib
import importlib.util
import json
from pathlib import Path
import struct
import xml.etree.ElementTree as ET
import zipfile

import ezdxf
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('project_tools',ROOT/'skills/interior-orchestrator/scripts/project_tools.py')
t=importlib.util.module_from_spec(spec);spec.loader.exec_module(t)

def require(condition,message):
    if not condition:raise ValueError(message)

def rect_bounds(entity):
    pts=list(entity.get_points('xy'));xs=[p[0] for p in pts];ys=[p[1] for p in pts]
    return tuple(round(x,4) for x in (min(xs),min(ys),max(xs)-min(xs),max(ys)-min(ys)))

def verify(p,artifacts_only=False,baseline_skills=False):
    m=t.load(p/'model/project.json');sha=t.sha(p/'model/project.json')
    model_check=t.check(m);require(model_check['passed'],'Project model failed')
    require((p/'visuals/render-input.json').read_bytes()==(p/'model/project.json').read_bytes(),'Stale render input')
    render=t.load(p/'visuals/render-manifest.json')
    require(render['model_sha256']==sha and render['revision']==m['revision'],'Render source hash differs')
    views=[]
    for row in render['views']:
        path=p/'visuals'/row['file']
        with Image.open(path) as im:
            require(im.size==(1600,1060),'Unexpected render size: '+row['file']);im.verify()
        views.append(row['file'])
    require(len(views)==3,'Expected three rendered views')
    blend=p/'visuals/qimu-home.blend';header=blend.read_bytes()[:7]
    require(header==b'BLENDER' or header[:4]==bytes.fromhex('28b52ffd'),'Unexpected Blender file format')
    blend_check=t.load(p/'state/blend-check.json')
    require(blend_check['passed'] and blend_check['sha256']==t.sha(blend) and blend_check['objects']>0,'Blender reopen evidence differs')
    raw=(p/'visuals/qimu-home.glb').read_bytes();magic,version,length=struct.unpack_from('<4sII',raw)
    require(magic==b'glTF' and version==2 and length==len(raw),'Invalid GLB header')
    chunk_length,chunk_type=struct.unpack_from('<II',raw,12)
    require(chunk_type==0x4E4F534A,'Missing GLB JSON chunk')
    gltf=json.loads(raw[20:20+chunk_length]);require(gltf.get('meshes') and gltf.get('scenes'),'Empty GLB')

    svg=[]
    for i in range(1,7):
        path=p/f'drawings/ID-{i:02d}.svg';doc=ET.parse(path)
        require(doc.getroot().tag.endswith('svg'),'Invalid SVG');svg.append(path.name)
    doc=ezdxf.readfile(p/'drawings/interior-set.dxf');audit=doc.audit();ms=doc.modelspace()
    require(not audit.errors and not audit.fixes,'DXF audit reported errors or fixes')
    require(doc.header['$INSUNITS']==4 and doc.units==4,'DXF is not in millimeters')
    by_layer=Counter(e.dxf.layer for e in ms);types=Counter(e.dxftype() for e in ms)
    for index in range(5):
        code=f'ID-{index+1:02d}';ox=(index%3)*17000;oy=-(index//3)*12500
        floor=ms.query(f'LWPOLYLINE[layer=="{code}_FLOOR"]')
        require(len(floor)==len(m['rooms']),'DXF room count differs: '+code)
        actual={rect_bounds(e) for e in floor}
        expected={tuple(round(v,4) for v in (r['rect'][0]+ox,r['rect'][1]+oy,*r['rect'][2:])) for r in m['rooms']}
        require(actual==expected,'DXF room dimensions differ: '+code)
        lines=ms.query(f'LINE[layer=="{code}_DIMENSION"]')
        require(any(abs(e.dxf.end.x-e.dxf.start.x-12000)<.01 for e in lines),'Missing 12000mm dimension')
        require(any(abs(e.dxf.end.y-e.dxf.start.y-8000)<.01 for e in lines),'Missing 8000mm dimension')
    actual={rect_bounds(e) for e in ms.query('LWPOLYLINE[layer=="ID-01_FURNITURE"]')}
    require(all(tuple(f['rect']) in actual for f in m['furniture']),'DXF furniture dimensions differ')
    require(all(e.closed for e in ms.query('LWPOLYLINE')),'Open rectangle profile in DXF')
    labels={e.dxf.text for e in ms.query('TEXT')}
    require(all(o['id']+' '+str(o['width']) in labels for o in m['openings']),'Missing opening dimension label')
    cabinet=ms.query('LWPOLYLINE[layer=="ID-06_JOINERY"]')
    require(any(rect_bounds(e)[2:]==(2400,2400) for e in cabinet),'Missing J01 elevation dimensions')

    pdf=PdfReader(p/'documents/栖木之家_方案与深化图集.pdf')
    require(len(pdf.pages)==12,'Expected twelve PDF pages')
    texts=[page.extract_text() or '' for page in pdf.pages]
    require(all(abs(float(page.mediabox.width)-1190.5512)<.1 and abs(float(page.mediabox.height)-841.8898)<.1 for page in pdf.pages),'PDF is not A3 landscape')
    require('栖木之家' in texts[0] and '287,907.60' in texts[4],'PDF Chinese text or final budget differs')
    require(all(f'ID-{i:02d}' in texts[5+i] for i in range(1,7)),'PDF drawing identifiers differ')
    embedded=False
    for page in pdf.pages:
        for ref in page['/Resources']['/Font'].values():
            font=ref.get_object();descriptor=font.get('/FontDescriptor')
            if descriptor and '/FontFile2' in descriptor.get_object():embedded=True
    require(embedded,'Chinese TrueType font is not embedded')

    workbook=p/'schedules/栖木之家_预算.xlsx'
    cached=load_workbook(workbook,read_only=True,data_only=True)
    formulas=load_workbook(workbook,read_only=True,data_only=False)
    require(cached.sheetnames==['预算总览','工程量计价'],'Unexpected workbook sheets')
    budget=t.totals(m);rows=t.budget_rows(m)
    for cell,key in [('C19','base'),('C20','reserve'),('C21','total'),('C22','remaining')]:
        require(round(cached['预算总览'][cell].value*100)==round(budget[key]*100),'Workbook amount differs: '+cell)
    for i,row in enumerate(rows,5):
        sheet=cached['工程量计价'];require(sheet[f'A{i}'].value==row['id'],'Budget row order differs')
        require(round(sheet[f'I{i}'].value*100)==round(row['amount']*100),'Budget line amount differs: '+row['id'])
        require(formulas['工程量计价'][f'I{i}'].data_type=='f','Budget amount is not a formula')
    cached.close();formulas.close()
    errors=(p/'state/workbook-error-scan.txt').read_text()
    require('matched 0 entries' in errors,'Workbook formula-error scan is not empty')
    issues=t.load(p/'state/issue-register.json')
    require(len(issues)==3 and all(x['status'].startswith('closed_') and (p/x['recheck_evidence']).is_file() for x in issues),'Design issues lack closure evidence')

    latest={}
    if not artifacts_only:
        for row in t.load(p/'state/execution-log.json'):latest[row['skill']]=row
        names={x.name for x in (ROOT/'skills').glob('interior-*')}
        require(set(latest)==names and len(names)==15,'Execution log does not cover all fifteen skills')
        for name,row in latest.items():
            if baseline_skills:
                with zipfile.ZipFile(ROOT/'releases/baseline-v1-tools.zip') as archive:expected_skill_sha=hashlib.sha256(archive.read('skills/'+name+'/SKILL.md')).hexdigest()
            else:expected_skill_sha=t.sha(ROOT/'skills'/name/'SKILL.md')
            require(row['skill_sha256']==expected_skill_sha,'Skill hash differs: '+name)
            require(row['input_sha256']==sha,'Latest stage used an older model: '+name)
            for out in row['output_files']:
                require((p/out['path']).is_file() and t.sha(p/out['path'])==out['sha256'],'Output hash differs: '+name+'/'+out['path'])
        manifest=t.load(p/'state/delivery-manifest.json')
        for out in manifest['files']:
            require(t.sha(p/out['path'])==out['sha256'],'Delivery manifest differs: '+out['path'])
    return {'passed':True,'revision':m['revision'],'model_sha256':sha,'model':model_check,
        'dxf':{'audit_errors':len(audit.errors),'audit_fixes':len(audit.fixes),'units':'mm','entity_counts':dict(types),'layer_count':len(by_layer),'room_rectangles_checked':35,'furniture_rectangles_checked':len(m['furniture']),'opening_labels_checked':len(m['openings']),'total_dimensions_checked':10},
        'svg_files':svg,'pdf':{'pages':len(pdf.pages),'format':'A3 landscape','chinese_text_extracted':True,'chinese_font_embedded':embedded},
        'model_exchange':{'blend_reopened_in_blender':True,'blend_objects':blend_check['objects'],'glb_version':version,'glb_meshes':len(gltf['meshes'])},
        'renders':{'input_hash_matches':True,'files':views},'budget':{'detail_rows':len(rows),'all_amounts_equal_to_cent':True,**budget},
        'closed_design_issues':len(issues),'latest_skill_records_checked':len(latest),'artifacts_only':artifacts_only,'baseline_skills':baseline_skills}

def main():
    ap=argparse.ArgumentParser();ap.add_argument('project',nargs='?',type=Path,default=ROOT/'projects/DEMO-APT-096');ap.add_argument('--artifacts-only',action='store_true');ap.add_argument('--baseline-skills',action='store_true');ap.add_argument('--out',type=Path);a=ap.parse_args();p=a.project.resolve()
    try:r=verify(p,a.artifacts_only,a.baseline_skills)
    except Exception as exc:r={'passed':False,'error':str(exc)}
    t.dump(a.out or p/'state/artifact-check.json',r)
    print(json.dumps(r,ensure_ascii=False,indent=2));return 0 if r['passed'] else 1

if __name__=='__main__':raise SystemExit(main())
